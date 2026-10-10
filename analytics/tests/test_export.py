import csv
import io
from datetime import date, datetime, timedelta
from io import StringIO
from zoneinfo import ZoneInfo

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from analytics import export
from analytics.models import SpikeAlert
from appeals.models import Appeal, Reason

MSK = ZoneInfo("Europe/Moscow")
PERIOD = {"date_from": "2026-09-11", "date_to": "2026-09-20"}
BOM = b"\xef\xbb\xbf"


def read_csv(response):
    content = response.content
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))


def read_xlsx(response):
    return load_workbook(io.BytesIO(response.content))


class ExportTests(TestCase):
    """Предыдущий период 1-10 сентября: Оплата 10. Текущий 11-20 сентября: Оплата 30, Доставка 12, без категории 3."""

    @classmethod
    def setUpTestData(cls):
        call_command("load_dictionary", stdout=StringIO())
        call_command("setup_roles", "--demo-users", stdout=StringIO())
        payment = Reason.objects.get(name="Двойное списание")
        delivery = Reason.objects.get(name="Задержка доставки")
        appeals = []

        def add(reason, count, first_day):
            for i in range(count):
                appeal = Appeal(
                    text="Списали дважды", created_at=datetime(2026, 9, first_day + i % 10, 10, i, tzinfo=MSK),
                    author_name="Иван Петров", author_email="ivan.petrov@mail.ru", category_source="none",
                )
                if reason:
                    appeal.category, appeal.reason, appeal.category_source = reason.category, reason, "auto"
                appeals.append(appeal)

        add(payment, 10, 1)
        add(payment, 30, 11)
        add(delivery, 12, 11)
        add(None, 3, 11)
        Appeal.objects.bulk_create(appeals)
        cls.formula = Appeal.objects.create(
            text='=HYPERLINK("http://example.com","жми")', created_at=datetime(2026, 9, 12, 12, tzinfo=MSK),
            author_name="Анна", author_email="anna@yandex.ru",
        )
        SpikeAlert.objects.create(
            category=payment.category, reason=payment, day=date(2026, 9, 15), count=9, baseline=1.5, ratio=6
        )

    def get(self, username, name, params):
        self.client.force_login(User.objects.get(username=username))
        return self.client.get(reverse(name), params)

    def report(self, **params):
        return self.get("analyst", "analytics:report_export", {**PERIOD, **params})

    def test_report_xlsx_has_five_sheets(self):
        response = self.report()
        self.assertEqual(response["Content-Type"], export.CONTENT_TYPES["xlsx"])
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="report_2026-09-11_2026-09-20.xlsx"'
        )
        book = read_xlsx(response)
        self.assertEqual(book.sheetnames, ["Сводка", "Категории", "Причины", "По дням", "Всплески"])

    def test_report_xlsx_numbers(self):
        book = read_xlsx(self.report())
        summary = {row[0]: row[1:] for row in book["Сводка"].iter_rows(values_only=True) if row[0]}
        self.assertEqual(summary["Период"][0], "11.09.2026 - 20.09.2026")
        self.assertEqual(summary["Обращений"][:2], (46, 10))
        self.assertAlmostEqual(summary["Обращений"][2], 3.6)
        categories = list(book["Категории"].iter_rows(values_only=True))
        self.assertEqual(categories[1][:4], ("Оплата", 30, 30 / 46, 10))
        self.assertAlmostEqual(categories[1][4], 2.0)
        self.assertEqual(book["Категории"]["E2"].number_format, "+0.0%;-0.0%;0.0%")
        daily = list(book["По дням"].iter_rows(values_only=True))
        self.assertEqual(daily[0][:2], ("День", "Всего"))
        self.assertEqual(len(daily), 11)
        self.assertEqual(sum(row[1] for row in daily[1:]), 46)
        alerts = list(book["Всплески"].iter_rows(values_only=True))
        self.assertEqual(alerts[1][1:4], ("Оплата", "Двойное списание", 9))

    def test_report_csv_sections(self):
        for section in export.SECTIONS:
            with self.subTest(section=section):
                response = self.report(format="csv", section=section)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.content.startswith(BOM))
                self.assertIn(f"_{section}.csv", response["Content-Disposition"])
        rows = read_csv(self.report(format="csv", section="categories"))
        self.assertEqual(rows[0], ["Категория", "Обращений", "Доля", "Было", "Изменение", "Заметное изменение"])
        self.assertEqual(rows[1], ["Оплата", "30", "65,2%", "10", "+200,0%", "да"])

    def test_unknown_section(self):
        self.assertEqual(self.report(format="csv", section="secret").status_code, 404)

    def test_empty_period(self):
        book = read_xlsx(self.report(date_from="2020-01-01", date_to="2020-01-31"))
        self.assertEqual(len(book.sheetnames), 5)
        self.assertEqual(book["Категории"].max_row, 1)

    def test_appeals_masked_for_analyst(self):
        rows = read_csv(self.get("analyst", "appeals:export", {"category": "payment"}))
        self.assertEqual(rows[0][:5], ["Номер", "Поступило", "Канал", "Автор", "Email"])
        self.assertEqual(len(rows), 41)
        self.assertEqual(rows[1][3:5], ["И*** П***", "i***@mail.ru"])
        self.assertNotIn("ivan.petrov@mail.ru", self.get("analyst", "appeals:export", {}).content.decode("utf-8-sig"))

    def test_appeals_full_for_admin(self):
        rows = read_csv(self.get("admin", "appeals:export", {"category": "payment"}))
        self.assertEqual(rows[1][3:5], ["Иван Петров", "ivan.petrov@mail.ru"])

    def test_appeals_xlsx(self):
        response = self.get("analyst", "appeals:export", {"format": "xlsx", "order": "old"})
        sheet = read_xlsx(response)["Обращения"]
        self.assertEqual(sheet.max_row, Appeal.objects.count() + 1)
        first = [cell.value for cell in sheet[2]]
        self.assertEqual(first[1], datetime(2026, 9, 1, 10, 0))  # московское время без пояса
        self.assertEqual(first[4], "i***@mail.ru")

    def test_formula_is_not_executed(self):
        rows = read_csv(self.get("admin", "appeals:export", {"q": f"#{self.formula.pk}"}))
        self.assertEqual(rows[1][6], "'" + self.formula.text)
        sheet = read_xlsx(self.get("admin", "appeals:export", {"q": f"#{self.formula.pk}", "format": "xlsx"}))
        cell = sheet["Обращения"]["G2"]
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'="))

    def test_operator_forbidden(self):
        for name in ("appeals:export", "analytics:report_export"):
            with self.subTest(url=name):
                self.assertEqual(self.get("operator", name, PERIOD).status_code, 403)

    def test_buttons(self):
        analyst = self.get("analyst", "appeals:list", {"category": "payment"})
        self.assertContains(analyst, "Выгрузить CSV")
        self.assertContains(analyst, 'href="/appeals/export/?category=payment&amp;format=xlsx"')
        self.assertNotContains(self.get("operator", "appeals:list", {}), "Выгрузить CSV")
        report = self.get("analyst", "analytics:reports", PERIOD)
        self.assertContains(report, "Скачать XLSX")
        self.assertContains(report, "section=categories")


class ValueFormatTests(TestCase):
    def test_csv_values(self):
        self.assertEqual(export.csv_value(export.Share(0.2642)), "26,4%")
        self.assertEqual(export.csv_value(export.Growth(-0.21)), "-21,0%")
        self.assertEqual(export.csv_value(41.5), "41,5")
        self.assertEqual(export.csv_value(2.0), "2")
        self.assertEqual(export.csv_value(date(2026, 9, 1)), "01.09.2026")
        self.assertEqual(export.csv_value(None), "")
        self.assertEqual(export.csv_value("-5 баллов"), "'-5 баллов")

    def test_sections_without_data(self):
        # За пустой период в отчёте есть только сводка, остальные разделы выгружаются одними заголовками.
        for key in ("categories", "reasons", "daily", "alerts"):
            with self.subTest(section=key):
                headers, rows = export.BUILDERS[key]({})
                self.assertTrue(headers)
                self.assertEqual(rows, [])


class LargeExportTests(TestCase):
    def test_many_rows(self):
        Appeal.objects.bulk_create(
            [Appeal(text=f"обращение {i}", created_at=datetime(2026, 9, 1, tzinfo=MSK) + timedelta(minutes=i))
             for i in range(3000)]
        )
        user = User.objects.create_user("x")
        rows = list(csv.reader(io.StringIO(export.appeals_csv(Appeal.objects.all(), user).decode("utf-8-sig")),
                               delimiter=";"))
        self.assertEqual(len(rows), 3001)
