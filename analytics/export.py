"""Выгрузка отчёта и списка обращений в XLSX и CSV (Д5).

Отчёт выгружается из тех же данных, что показывает страница отчёта (reports.build), поэтому числа совпадают.
XLSX: по листу на раздел. CSV: один раздел на файл, UTF-8 с BOM и разделителем «;», так файл без настройки
открывается в русском Excel. Имя и email автора в списке обращений проходят через roles.author_for:
у аналитика они замаскированы так же, как на карточке.
"""

import csv
import io
from datetime import date, datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from appeals.roles import author_for

from . import stats
from .reports import change

SECTIONS = {
    "summary": "Сводка",
    "categories": "Категории",
    "reasons": "Причины",
    "daily": "По дням",
    "alerts": "Всплески",
}
# Ячейка, которая начинается с этих символов, в Excel станет формулой. Тексты обращений пишут клиенты.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
MAX_WIDTH = 60
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class Share(float):
    """Доля 0..1. В XLSX получает процентный формат, в CSV пишется как «26,4%»."""


class Growth(Share):
    """Изменение к прошлому периоду: процент со знаком, «+4,9%» и «-21,0%»."""


def safe_text(value):
    return "'" + value if value.startswith(FORMULA_START) else value


def csv_value(value):
    if value is None:
        return ""
    if isinstance(value, Growth):
        return f"{value * 100:+.1f}%".replace(".", ",")
    if isinstance(value, Share):
        return f"{value * 100:.1f}%".replace(".", ",")
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    if isinstance(value, datetime):
        return f"{timezone.localtime(value):%d.%m.%Y %H:%M}"
    if isinstance(value, date):
        return f"{value:%d.%m.%Y}"
    if isinstance(value, str):
        return safe_text(value)
    return value


def growth(value):
    return None if value is None else Growth(value)


def summary_section(report):
    now, before = report["summary"], report["previous"]
    rows = [
        ["Обращений", now.total, before.total],
        ["В среднем в день", now.per_day, before.per_day],
        ["Открытые: новые и в работе", now.open_count, before.open_count],
        ["Без категории", now.without_category, before.without_category],
        ["Доля без категории", Share(now.without_category_share), Share(before.without_category_share)],
        ["Медиана времени решения, ч", now.median_resolution_hours, before.median_resolution_hours],
    ]
    for row in rows:
        current, previous = row[1], row[2]
        if isinstance(current, Share) or current is None or previous is None:
            row.append(None)
        else:
            row.append(growth(change(current, previous)))
    return ["Показатель", "Период", "Предыдущий период", "Изменение"], rows


def categories_section(report):
    rows = [
        [r["name"], r["current_count"], Share(r["share"]), r["base_count"], growth(r["change"]),
         "да" if r["notable"] else "нет"]
        for r in report.get("categories", [])
    ]
    return ["Категория", "Обращений", "Доля", "Было", "Изменение", "Заметное изменение"], rows


def reasons_section(report):
    rows = [
        [r["rank"], r["name"], r["category"], r["count"], Share(r["share"]), r["prev_count"], growth(r["growth"]),
         Share(r["open_share"]), r["median_resolution_hours"]]
        for r in report.get("problems", [])
    ]
    headers = ["Место", "Причина", "Категория", "Обращений", "Доля", "Было", "Прирост", "Нерешённых",
               "Медиана решения, ч"]
    return headers, rows


def daily_section(report):
    table = report.get("daily")
    if table is None or table.empty:
        return ["День", "Всего"], []
    table = table[table.sum().sort_values(ascending=False).index]
    rows = [[day, int(values.sum()), *(int(v) for v in values)] for day, values in table.iterrows()]
    return ["День", "Всего", *table.columns], rows


def alerts_section(report):
    rows = [
        [a.day, a.category.name, a.reason.name if a.reason_id else "вся категория", a.count, a.baseline, a.ratio]
        for a in report.get("alerts", [])
    ]
    return ["День", "Категория", "Причина", "Обращений", "Обычно", "Рост, раз"], rows


BUILDERS = {
    "summary": summary_section,
    "categories": categories_section,
    "reasons": reasons_section,
    "daily": daily_section,
    "alerts": alerts_section,
}


def report_title(report, form):
    """Строки шапки листа «Сводка»: период, фильтры, время выгрузки."""
    data = form.cleaned_data
    filters = []
    for name in ("category", "channel"):
        value = data.get(name)
        if value:
            labels = dict(form.fields[name].choices)
            filters.append(f"{form.fields[name].label.lower()}: {labels.get(value, value)}")
    return [
        ["Аналитический отчёт"],
        ["Период", f"{report['start']:%d.%m.%Y} - {report['end']:%d.%m.%Y}"],
        ["Предыдущий период", f"{report['prev_start']:%d.%m.%Y} - {report['prev_end']:%d.%m.%Y}"],
        ["Фильтры", ", ".join(filters) or "нет"],
        ["Сформирован", f"{timezone.localtime():%d.%m.%Y %H:%M}"],
        [],
    ]


def xlsx_value(value):
    """Excel не хранит часовой пояс: время переводится в московское без пояса."""
    if isinstance(value, str):
        return safe_text(value)
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def number_format(value):
    if isinstance(value, Growth):
        return "+0.0%;-0.0%;0.0%"
    if isinstance(value, Share):
        return "0.0%"
    if isinstance(value, float):
        return "0.0#"
    if isinstance(value, datetime):
        return "DD.MM.YYYY HH:MM"
    if isinstance(value, date):
        return "DD.MM.YYYY"
    return None


def write_sheet(sheet, headers, rows, title_rows=()):
    for line in title_rows:
        sheet.append([xlsx_value(v) for v in line])
    if title_rows:
        sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
    header_row = len(title_rows) + 1
    sheet.append(list(headers))
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
    widths = [len(str(h)) for h in headers]
    for line in title_rows:
        for column, value in enumerate(line[: len(widths)]):
            widths[column] = max(widths[column], len(str(value)))
    # Номер строки считается вручную: sheet.max_row перебирает все ячейки и на тысячах строк тормозит.
    for row_number, values in enumerate(rows, start=header_row + 1):
        sheet.append([xlsx_value(v) for v in values])
        for column, value in enumerate(values, start=1):
            fmt = number_format(value)
            if fmt:
                sheet.cell(row=row_number, column=column).number_format = fmt
            widths[column - 1] = max(widths[column - 1], len(str(csv_value(value))))
    for column, width in enumerate(widths, start=1):
        # Ширина в Excel меряется в символах латиницы, кириллица и жирный заголовок шире примерно на пятую часть.
        sheet.column_dimensions[get_column_letter(column)].width = min(round(width * 1.2) + 2, MAX_WIDTH)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)


def to_xlsx(sheets):
    """sheets: список (название листа, заголовки, строки, строки шапки). Возвращает байты файла."""
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows, title_rows in sheets:
        write_sheet(book.create_sheet(title), headers, rows, title_rows)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_csv(headers, rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(headers)
    for values in rows:
        writer.writerow([csv_value(v) for v in values])
    return buffer.getvalue().encode("utf-8-sig")


def report_xlsx(report, form):
    sheets = []
    for key, title in SECTIONS.items():
        headers, rows = BUILDERS[key](report)
        sheets.append((title, headers, rows, report_title(report, form) if key == "summary" else ()))
    return to_xlsx(sheets)


def report_csv(report, section):
    return to_csv(*BUILDERS[section](report))


APPEAL_HEADERS = [
    "Номер", "Поступило", "Канал", "Автор", "Email", "Тема", "Текст", "Статус", "Категория", "Причина",
    "Источник категории", "Теги", "Исполнитель", "Решено",
]


def appeal_rows(queryset, user):
    """Строки выгрузки обращений. Имя и email в том виде, в каком их можно показать этому пользователю."""
    queryset = queryset.select_related("category", "reason", "assignee").prefetch_related("tags")
    for appeal in queryset.iterator(chunk_size=2000):
        name, email = author_for(user, appeal)
        assignee = ""
        if appeal.assignee_id:
            assignee = appeal.assignee.get_full_name() or appeal.assignee.username
        yield [
            appeal.pk,
            appeal.created_at,
            appeal.get_channel_display(),
            name,
            email,
            appeal.subject,
            appeal.text,
            appeal.get_status_display(),
            appeal.category.name if appeal.category_id else stats.NO_CATEGORY,
            appeal.reason.name if appeal.reason_id else "",
            appeal.get_category_source_display(),
            ", ".join(tag.name for tag in appeal.tags.all()),
            assignee,
            appeal.resolved_at,
        ]


def appeals_xlsx(queryset, user):
    return to_xlsx([("Обращения", APPEAL_HEADERS, list(appeal_rows(queryset, user)), ())])


def appeals_csv(queryset, user):
    return to_csv(APPEAL_HEADERS, appeal_rows(queryset, user))


def download(content, filename):
    """Ответ с файлом. Тип берётся по расширению имени."""
    response = HttpResponse(content, content_type=CONTENT_TYPES[filename.rsplit(".", 1)[1]])
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
